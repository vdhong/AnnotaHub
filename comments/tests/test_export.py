"""
Test tính đúng đắn của chức năng xuất dữ liệu.

Trọng tâm là những thứ không thể phát hiện bằng mắt thường:
- BIO có phân biệt được hai cụm liền kề cùng nhãn hay không.
- Offset của cụm có trỏ đúng vào văn bản hay không.
- Giá trị "chưa xác định" có bị ghi thành "sai" hay không.
"""
import csv
import io
import json
import xml.etree.ElementTree as ET

from django.test import TestCase

from comments.export_service import (
    EXPORT_FORMATS,
    _comments_queryset,
    bio_tags,
    entity_spans,
    resolve_format,
    safe_filename,
    tri_state,
)
from comments.models import Token
from comments.services import annotation as annotation_service
from comments.tests.factories import (
    make_comment,
    make_label,
    make_link,
    make_project,
    make_project_label,
    make_user,
)

PASSWORD = 'test-pass-12345'


class TriStateTests(TestCase):
    """`is_meaningful = None` nghĩa là chưa xét, không phải 'vô nghĩa'."""

    def test_none_is_unknown_not_false(self):
        self.assertEqual(tri_state(None), 'unknown')
        self.assertEqual(tri_state(False), 'false')
        self.assertEqual(tri_state(True), 'true')


class BioTaggingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = make_user('bio_owner')
        cls.project = make_project(cls.owner, name='BIO')
        cls.toxic = make_project_label(cls.project, make_label(cls.owner, 'toxic'))
        cls.other = make_project_label(cls.project, make_label(cls.owner, 'spam'))
        cls.link = make_link(cls.project)

    # Văn bản dùng cho test được chọn sao cho tokenizer tách ổn định, nhưng vị
    # trí token vẫn được tra cứu động thay vì cứng hoá: pyvi gộp từ ghép tiếng
    # Việt (ví dụ "một hai" thành một token), nên mọi giả định về chỉ số cố định
    # đều dễ vỡ khi đổi tokenizer.
    TEXT = 'thằng ngu này nói linh tinh đồ khốn nạn'

    def _comment(self, text, comment_id):
        comment = make_comment(self.link, text=text, comment_id=comment_id)
        comment.ensure_token_inventory()
        return comment

    @staticmethod
    def _positions(comment, *words):
        """Vị trí token khớp với từng từ cho trước, theo thứ tự xuất hiện."""
        tokens = comment.display_tokens
        found = []
        for word in words:
            for token in tokens:
                if token['text'] == word and token['position'] not in found:
                    found.append(token['position'])
                    break
            else:
                raise AssertionError(
                    f'Không tìm thấy token {word!r} trong {[t["text"] for t in tokens]}'
                )
        return found

    def test_single_span_gets_b_then_i(self):
        comment = self._comment(self.TEXT, 'bio-1')
        first, second = self._positions(comment, 'thằng', 'ngu')
        annotation_service.set_token_span_label(
            comment, self.owner, first, second, self.toxic
        )
        tags = bio_tags(comment.display_tokens)
        self.assertEqual(tags[first], 'B-toxic')
        self.assertEqual(tags[second], 'I-toxic')
        # Mọi token còn lại không mang nhãn.
        for index, tag in enumerate(tags):
            if index not in (first, second):
                self.assertEqual(tag, 'O')

    def test_two_separate_spans_same_label_are_distinguished(self):
        """
        Điều không thể làm được với nhãn phẳng: hai cụm rời cùng nhãn phải ra
        B-…, I-… rồi lại B-…, I-… chứ không phải một chuỗi I- liên tục.
        """
        comment = self._comment(self.TEXT, 'bio-2')
        a1, a2 = self._positions(comment, 'thằng', 'ngu')
        b1, b2 = self._positions(comment, 'đồ', 'khốn nạn')

        annotation_service.set_token_span_label(comment, self.owner, a1, a2, self.toxic)
        annotation_service.set_token_span_label(comment, self.owner, b1, b2, self.toxic)

        tags = bio_tags(comment.display_tokens)
        self.assertEqual(tags[a1], 'B-toxic')
        self.assertEqual(tags[a2], 'I-toxic')
        # Cụm thứ hai phải bắt đầu lại bằng B-, không nối tiếp thành I-.
        self.assertEqual(tags[b1], 'B-toxic')
        self.assertEqual(tags[b2], 'I-toxic')

        spans = entity_spans(comment)
        self.assertEqual(len(spans), 2, f'Phải là 2 cụm riêng, nhận được {spans}')
        self.assertEqual(spans[0]['text'], 'thằng ngu')
        self.assertEqual(spans[1]['text'], 'đồ khốn nạn')

    def test_adjacent_separate_single_tokens_are_two_spans(self):
        """Hai token liền kề gán riêng lẻ vẫn phải là hai cụm khác nhau."""
        comment = self._comment(self.TEXT, 'bio-3')
        first, second = self._positions(comment, 'thằng', 'ngu')

        annotation_service.set_token_label(comment, self.owner, first, self.toxic)
        annotation_service.set_token_label(comment, self.owner, second, self.toxic)

        tags = bio_tags(comment.display_tokens)
        self.assertEqual(tags[first], 'B-toxic')
        self.assertEqual(tags[second], 'B-toxic')
        self.assertEqual(len(entity_spans(comment)), 2)

    def test_different_labels_never_merge(self):
        comment = self._comment(self.TEXT, 'bio-4')
        first, second = self._positions(comment, 'thằng', 'ngu')
        annotation_service.set_token_label(comment, self.owner, first, self.toxic)
        annotation_service.set_token_label(comment, self.owner, second, self.other)
        tags = bio_tags(comment.display_tokens)
        self.assertEqual(tags[first], 'B-toxic')
        self.assertEqual(tags[second], 'B-spam')

    def test_legacy_tokens_without_span_group_merge_by_run(self):
        """Dữ liệu cũ (span_group=NULL) gộp theo chuỗi liền kề cùng nhãn."""
        comment = self._comment(self.TEXT, 'bio-5')
        first, second = self._positions(comment, 'thằng', 'ngu')
        Token.objects.filter(
            comment=comment, position__in=(first, second)
        ).update(manual_label=self.toxic, span_group=None)

        tags = bio_tags(comment.display_tokens)
        self.assertEqual(tags[first], 'B-toxic')
        self.assertEqual(tags[second], 'I-toxic')

    def test_span_offsets_point_at_real_text(self):
        comment = self._comment(self.TEXT, 'bio-6')
        first, second = self._positions(comment, 'nói', 'linh tinh')
        annotation_service.set_token_span_label(
            comment, self.owner, first, second, self.toxic
        )
        spans = entity_spans(comment)
        self.assertTrue(spans)
        for span in spans:
            self.assertEqual(comment.text[span['start']:span['end']], span['text'])
            self.assertTrue(span['text'].strip())


class ExportFormatTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = make_user('exp_owner')
        cls.project = make_project(cls.owner, name='Xuất dữ liệu')
        cls.toxic = make_project_label(cls.project, make_label(cls.owner, 'toxic'))
        cls.link = make_link(cls.project)

        cls.labelled = make_comment(cls.link, text='thằng ngu này im đi',
                                    comment_id='exp-1')
        cls.labelled.ensure_token_inventory()
        annotation_service.set_token_span_label(cls.labelled, cls.owner, 0, 1, cls.toxic)
        annotation_service.set_comment_label(cls.labelled, cls.owner, cls.toxic)

        # Bình luận chưa ai xét: is_meaningful = None
        cls.untouched = make_comment(cls.link, text='bình thường thôi',
                                     comment_id='exp-2')

    def _queryset(self, review_filter='all'):
        return _comments_queryset(self.project, self.link, 'all',
                                  review_filter=review_filter)

    def _render(self, key, review_filter='all'):
        meta = EXPORT_FORMATS[key]
        queryset = self._queryset(review_filter)
        if meta.streaming:
            return ''.join(meta.builder(queryset, project=self.project))
        return meta.builder(queryset, project=self.project)

    def test_every_format_produces_output(self):
        for key in EXPORT_FORMATS:
            with self.subTest(format=key):
                payload = self._render(key)
                self.assertTrue(len(payload) > 0)

    def test_json_formats_are_valid_json(self):
        for key in ('json_sentence', 'json_token', 'spacy_json', 'label_studio_json'):
            with self.subTest(format=key):
                json.loads(self._render(key))

    def test_jsonl_formats_are_line_delimited_json(self):
        for key in ('jsonl', 'hf_jsonl', 'doccano_jsonl', 'json_llm'):
            with self.subTest(format=key):
                lines = [line for line in self._render(key).splitlines() if line.strip()]
                self.assertTrue(lines)
                for line in lines:
                    json.loads(line)

    def test_xml_is_well_formed(self):
        ET.fromstring(self._render('xml'))

    def test_conll_is_tab_separated_plain_text(self):
        """CoNLL thật: token <TAB> nhãn, dòng trống ngăn câu."""
        payload = self._render('conll')
        body = [line for line in payload.splitlines()
                if line and not line.startswith('#')]
        self.assertTrue(body)
        for line in body:
            parts = line.split('\t')
            self.assertEqual(len(parts), 2, f'Dòng không đúng 2 cột: {line!r}')
            self.assertTrue(parts[1] == 'O' or parts[1][:2] in ('B-', 'I-'))
        self.assertIn('\n\n', payload)

    def test_legacy_xml_conll_alias_still_works(self):
        self.assertEqual(resolve_format('xml_conll'), 'xml')

    def test_hf_jsonl_has_label_map_and_matching_tags(self):
        lines = [
            json.loads(line)
            for line in self._render('hf_jsonl').splitlines() if line.strip()
        ]
        meta = lines[0]
        self.assertTrue(meta.get('__meta__'))
        self.assertEqual(meta['label_names'][0], 'O')
        self.assertIn('B-toxic', meta['label_names'])

        for row in lines[1:]:
            self.assertEqual(len(row['tokens']), len(row['ner_tags']))
            for tag_id, tag in zip(row['ner_tags'], row['bio_tags'], strict=False):
                self.assertEqual(meta['label_names'][tag_id], tag)

    def test_spacy_entities_offsets_match_text(self):
        for text, meta in json.loads(self._render('spacy_json')):
            for start, end, _label in meta['entities']:
                self.assertTrue(text[start:end].strip())

    def test_doccano_offsets_match_text(self):
        for line in self._render('doccano_jsonl').splitlines():
            row = json.loads(line)
            for start, end, _label in row['label']:
                self.assertTrue(row['text'][start:end].strip())

    def test_csv_marks_unknown_meaningfulness(self):
        payload = self._render('csv_sentence').lstrip('﻿')
        rows = list(csv.DictReader(io.StringIO(payload)))
        by_id = {row['id']: row for row in rows}
        self.assertEqual(by_id['exp-2']['is_meaningful'], 'unknown')
        self.assertEqual(by_id['exp-1']['is_meaningful'], 'true')

    def test_csv_token_includes_bio_column(self):
        payload = self._render('csv_token').lstrip('﻿')
        rows = list(csv.DictReader(io.StringIO(payload)))
        self.assertIn('bio', rows[0])
        self.assertTrue(any(row['bio'].startswith('B-') for row in rows))

    def test_csv_spans_only_lists_labelled_spans(self):
        payload = self._render('csv_spans').lstrip('﻿')
        rows = list(csv.DictReader(io.StringIO(payload)))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['span_text'], 'thằng ngu')

    def test_xlsx_opens_as_workbook(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('openpyxl chưa được cài')
        workbook = load_workbook(io.BytesIO(self._render('xlsx')))
        self.assertIn('Câu', workbook.sheetnames)
        self.assertIn('Cụm đã gán nhãn', workbook.sheetnames)


class ReviewFilterTests(TestCase):
    """Bộ lọc phạm vi: quyết định dữ liệu nào được đưa vào dataset công bố."""

    @classmethod
    def setUpTestData(cls):
        cls.owner = make_user('rev_owner')
        cls.project = make_project(cls.owner, name='Phạm vi')
        cls.label = make_project_label(cls.project, make_label(cls.owner, 'toxic'))
        cls.link = make_link(cls.project)

        cls.gold = make_comment(cls.link, text='đã chốt', comment_id='rev-gold')
        annotation_service.set_comment_label(cls.gold, cls.owner, cls.label)

        cls.ai_only = make_comment(cls.link, text='ai gán', comment_id='rev-ai',
                                   ai_label=cls.label, ai_processed=True)
        cls.blank = make_comment(cls.link, text='chưa ai xét', comment_id='rev-blank')

    def _ids(self, review_filter):
        queryset = _comments_queryset(self.project, self.link, 'all',
                                      review_filter=review_filter)
        return set(queryset.values_list('youtube_comment_id', flat=True))

    def test_all_includes_everything(self):
        self.assertEqual(self._ids('all'), {'rev-gold', 'rev-ai', 'rev-blank'})

    def test_labelled_excludes_untouched(self):
        self.assertEqual(self._ids('labelled'), {'rev-gold', 'rev-ai'})

    def test_human_excludes_ai_only(self):
        self.assertEqual(self._ids('human'), {'rev-gold'})

    def test_gold_only_returns_settled_labels(self):
        self.assertEqual(self._ids('gold'), {'rev-gold'})

    def test_export_api_accepts_review_parameter(self):
        self.client.login(username='rev_owner', password=PASSWORD)
        response = self.client.post(
            f'/api/links/{self.link.id}/export/',
            data={'format': 'conll', 'review': 'gold'},
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        body = b''.join(response.streaming_content).decode()
        self.assertIn('rev-gold', body)
        self.assertNotIn('rev-blank', body)


class FilenameTests(TestCase):
    def test_dangerous_project_name_is_sanitised(self):
        name = safe_filename('Dự án "nguy hiểm"\nX')
        self.assertNotIn('"', name)
        self.assertNotIn('\n', name)
        self.assertTrue(name)


class FormatMetadataTests(TestCase):
    """Mỗi định dạng phải có nhãn, nhóm, mô tả và ví dụ để hiển thị trên UI."""

    def test_every_format_has_complete_metadata(self):
        for key, meta in EXPORT_FORMATS.items():
            with self.subTest(format=key):
                self.assertTrue(str(meta.label).strip(), f'{key}: thiếu label')
                self.assertTrue(str(meta.group).strip(), f'{key}: thiếu group')
                self.assertTrue(str(meta.description).strip(),
                                f'{key}: thiếu description')
                self.assertTrue(str(meta.sample).strip(), f'{key}: thiếu sample')
                self.assertTrue(meta.extension.startswith('.'))
                self.assertTrue(callable(meta.builder))

    def test_format_choices_exposes_description_and_sample(self):
        from comments.export_service import format_choices

        choices = format_choices()
        self.assertEqual(len(choices), len(EXPORT_FORMATS))
        for choice in choices:
            for field in ('key', 'label', 'group', 'extension', 'streaming',
                          'description', 'sample'):
                self.assertIn(field, choice)

    def test_format_groups_keeps_declaration_order(self):
        from comments.export_service import format_groups

        grouped = format_groups()
        self.assertTrue(grouped)
        total = sum(len(items) for items in grouped.values())
        self.assertEqual(total, len(EXPORT_FORMATS))

    def test_labels_are_translated(self):
        """Nhãn dùng gettext_lazy nên đổi theo ngôn ngữ đang hoạt động."""
        from django.utils import translation

        with translation.override('vi'):
            vi_labels = [str(m.label) for m in EXPORT_FORMATS.values()]
            vi_groups = {str(m.group) for m in EXPORT_FORMATS.values()}
        with translation.override('en'):
            en_labels = [str(m.label) for m in EXPORT_FORMATS.values()]
            en_groups = {str(m.group) for m in EXPORT_FORMATS.values()}

        self.assertNotEqual(vi_labels, en_labels,
                            'Nhãn định dạng phải khác nhau giữa hai ngôn ngữ.')
        self.assertNotEqual(vi_groups, en_groups,
                            'Tên nhóm phải khác nhau giữa hai ngôn ngữ.')


class ExportPageI18nTests(TestCase):
    """Trang xuất dữ liệu phải dịch được toàn bộ, kể cả hộp select và mô tả."""

    @classmethod
    def setUpTestData(cls):
        cls.owner = make_user('exp_i18n')
        cls.project = make_project(cls.owner, name='Xuất i18n')
        make_project_label(cls.project, make_label(cls.owner, 'toxic'))
        cls.link = make_link(cls.project)
        make_comment(cls.link, comment_id='ei-1')

    def setUp(self):
        self.client.login(username='exp_i18n', password=PASSWORD)

    def _body(self, lang):
        response = self.client.get(f'/{lang}/projects/{self.project.id}/export/')
        self.assertEqual(response.status_code, 200)
        return response.content.decode()

    def test_all_formats_appear_in_select(self):
        body = self._body('vi')
        for key in EXPORT_FORMATS:
            self.assertIn(f'value="{key}"', body, f'Thiếu {key} trong hộp select')

    def test_all_formats_have_a_description_panel(self):
        for lang in ('vi', 'en'):
            body = self._body(lang)
            for key in EXPORT_FORMATS:
                self.assertIn(f'id="doc-{key}"', body,
                              f'/{lang}/: thiếu mô tả cho {key}')

    def test_select_labels_are_translated(self):
        vi_body = self._body('vi')
        en_body = self._body('en')

        self.assertIn('JSON — cấp câu', vi_body)
        self.assertIn('JSON — sentence level', en_body)
        self.assertIn('Chuẩn NLP', vi_body)
        self.assertIn('NLP standards', en_body)

    def test_english_page_has_no_vietnamese_format_labels(self):
        """Nhãn tiếng Việt lọt vào trang tiếng Anh nghĩa là thiếu bản dịch."""
        body = self._body('en')
        for text in ('JSON — cấp câu', 'CSV — cấp token (kèm BIO)',
                     'Bảng tính', 'Huấn luyện mô hình',
                     'CSV — chỉ các cụm đã gán nhãn'):
            self.assertNotIn(text, body, f'Còn sót tiếng Việt: {text!r}')

    def test_descriptions_are_translated(self):
        self.assertIn('Định dạng chuẩn của bài toán gán nhãn chuỗi', self._body('vi'))
        self.assertIn('The standard format for sequence labelling', self._body('en'))

    def test_output_file_hint_is_translated(self):
        self.assertIn('Tệp kết quả', self._body('vi'))
        self.assertIn('Output file', self._body('en'))
        self.assertNotIn('Tệp kết quả', self._body('en'))

    def test_api_export_formats_returns_plain_strings(self):
        """Chuỗi lazy phải được ép sang str để JSON hoá được."""
        import json

        response = self.client.get('/api/export-formats/')
        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertEqual(len(payload['formats']), len(EXPORT_FORMATS))
        for item in payload['formats']:
            for field in ('label', 'group', 'description', 'sample'):
                self.assertIsInstance(item[field], str)
                self.assertTrue(item[field].strip())
